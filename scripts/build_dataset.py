"""Build static review data from the revised photos and the frozen AI workbook."""
from __future__ import annotations
from pathlib import Path
from collections import Counter,defaultdict
import argparse
import hashlib
import json
import re
import shutil
import sys
import openpyxl

DEFAULT_PROJECT=Path(__file__).resolve().parents[1]
HAZARD_TYPES=['墜落、滾落','跌倒','衝撞','物體飛落','物體倒塌、崩塌','被撞','被夾、被捲','被切、割、擦傷','踩踏 (踏穿)','溺斃','與高溫、低溫之接觸','與有害物等之接觸','感電','爆炸','物體破裂','火災','不當動作','其他','無法歸類者']
LAW_PCODE={
 '職業安全衛生設施規則':'N0060009',
 '營造安全衛生設施標準':'N0060014',
 '職業安全衛生法':'N0060001',
 '缺氧症預防規則':'N0060020',
 '起重升降機具安全規則':'N0060013',
 '職業安全衛生管理辦法':'N0060027',
 '機械設備器具安全標準':'N0060034',
}
GUIDE_URLS={'高氣溫作業熱危害預防指引':'https://www.osha.gov.tw/1106/1196/10141/18360/18362/'}
LAW_ALIASES={'營造安全衛生設施規則':'營造安全衛生設施標準'}
LAW_NAME=r'[\u4e00-\u9fff]{2,40}?(?:規則|標準|辦法|指引|法)'
ARTICLE=r'第\s*(?P<main>\d+)\s*(?:(?P<dash>[-－之])\s*(?P<dashn>\d+)\s*)?條(?:\s*之\s*(?P<after>\d+))?(?:\s*第\s*\d+\s*[項款])?'
CITATION=re.compile(rf'(?:(?P<law>{LAW_NAME})\s*)?{ARTICLE}')
def digest(path:Path)->str:
 return hashlib.sha256(path.read_bytes()).hexdigest()
def write_json(path:Path,value)->None:
 path.parent.mkdir(parents=True,exist_ok=True)
 path.write_text(json.dumps(value,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
def canonical_law(name:str)->str:
 name=re.sub(r'\s+','',str(name or ''))
 return LAW_ALIASES.get(name,name)
def article_from_match(match)->str:
 if not match or not match.group('main'):return ''
 sub=match.group('after') or match.group('dashn')
 return f"{match.group('main')}-{sub}" if sub else match.group('main')
def article_number(value:str)->str:
 return article_from_match(CITATION.search(f'第{value}條' if '條' not in str(value) else str(value)))
def detail_identity(detail:dict)->tuple[str,str]:
 raw_name=str(detail.get('法規名稱') or '').strip()
 glued=re.match(rf'^(?P<law>{LAW_NAME})\s*{ARTICLE}$',raw_name)
 if glued:return canonical_law(glued.group('law')),article_from_match(glued) or article_number(detail.get('條號'))
 return canonical_law(raw_name),article_number(detail.get('條號'))
def official_link(law:str,article:str)->tuple[str,bool]:
 if law in GUIDE_URLS:return GUIDE_URLS[law],False
 pcode=LAW_PCODE.get(law)
 if pcode and article:return f'https://law.moj.gov.tw/LawClass/LawSingle.aspx?pcode={pcode}&flno={article}',True
 if pcode:return f'https://law.moj.gov.tw/LawClass/LawAll.aspx?pcode={pcode}',False
 return 'https://law.moj.gov.tw/',False
def citation_present(text:str)->bool:
 return bool(CITATION.search(str(text or '')))
def catalog_from_details(details)->dict:
 counts=defaultdict(Counter)
 for detail in details:
  law,article=detail_identity(detail)
  body=str(detail.get('完整法規內文') or '').strip()
  if law and article and body:counts[(law,article)][body]+=1
 return {key:counter.most_common(1)[0][0] for key,counter in counts.items()}
def unique_law(article:str,photo:dict,catalog:dict)->str:
 found={law for law,art in photo if art==article}
 if len(found)==1:return next(iter(found))
 found={law for law,art in catalog if art==article}
 if len(found)==1:return next(iter(found))
 return ''
def split_refs(text:str,details:list,catalog:dict|None=None)->tuple[list,list]:
 """Pair citation strings with article text already stored in the AI workbook.
 Abbreviated follow-on numbers keep the preceding statute, and 第N條之M stays one article."""
 catalog=catalog or {}
 text=str(text or '')
 photo={}
 for detail in details or []:
  law,article=detail_identity(detail)
  body=str(detail.get('完整法規內文') or '').strip()
  if not law or not article or not body:continue
  photo.setdefault((law,article),[])
  if body not in photo[(law,article)]:photo[(law,article)].append(body)
 refs=[];seen=set();current='';unresolved=[]
 for match in CITATION.finditer(text):
  named=canonical_law(match.group('law') or '')
  if named:current=named
  article=article_from_match(match)
  law=current or unique_law(article,photo,catalog)
  if not law:
   unresolved.append(match.group(0));continue
  key=(law,article)
  if key in seen:continue
  seen.add(key)
  texts=photo.get(key) or ([catalog[key]] if key in catalog else [])
  url,specific=official_link(law,article)
  refs.append({'id':f'l{len(refs)+1:02d}','label':f'{law}第{article}條','law_name':law,'article':article,
   'ai_texts':texts,'official_url':url,'official_link_is_specific':specific})
 return refs,unresolved

A11_LOAD_WEIGHT=0.88
READING_SPREAD_LIMIT=800
def reading_load(record:dict)->int:
 """Characters an expert has to judge: description, rationale, law text and cases."""
 n=len(record.get('description') or '')
 for scene in record.get('vlm') or []:
  n+=len((scene.get('危害情境推演') if isinstance(scene,dict) else '') or '')
 for hazard in record.get('hazards') or []:
  n+=len(hazard.get('law_text') or '')
  for law in hazard.get('laws') or []:
   for text in law.get('ai_texts') or []:n+=len(text or '')
  for case in hazard.get('cases') or []:n+=len(case.get('text') or '')
 return n
def assign_photos(index:list)->dict:
 """Disjoint groups of 30. Domain quotas stay 20/10, except A11 at 19/11.
 A01–A10 share judged text evenly; A11 is filled toward 88% of that load.
 Higher-hazard photos move off A11, then A01–A10 swap until their hazard counts
 differ by at most 10, while judged text stays within 800 characters."""
 experts=[f'A{i:02d}' for i in range(1,12)];others=experts[:-1]
 weight={e:(A11_LOAD_WEIGHT if e=='A11' else 1.0) for e in experts}
 caps={'營造':{e:(19 if e=='A11' else 20) for e in experts},'製造':{e:(11 if e=='A11' else 10) for e in experts}}
 by={p['id']:p for p in index}
 loads={e:0 for e in experts};counts={d:{e:0 for e in experts} for d in caps};groups={e:[] for e in experts}
 for photo in sorted(index,key=lambda p:(-p['reading_load'],p['id'])):
  eligible=[e for e in experts if counts[photo['domain']][e]<caps[photo['domain']][e]]
  chosen=min(eligible,key=lambda e:(loads[e]/weight[e],counts[photo['domain']][e],e))
  groups[chosen].append(photo['id']);loads[chosen]+=photo['reading_load'];counts[photo['domain']][chosen]+=1
 def judged(expert):return sum(by[i]['reading_load'] for i in groups[expert])
 def hazards(expert):return sum(by[i]['hazard_count'] for i in groups[expert])
 for _ in range(8):
  moved=False
  for pid in sorted(list(groups['A11']),key=lambda i:(-by[i]['hazard_count'],by[i]['reading_load'],i)):
   if hazards('A11')<=min(hazards(e) for e in others) and judged('A11')<min(judged(e) for e in others):break
   best=None
   dom=by[pid]['domain']
   current={e:judged(e) for e in others}
   a11=judged('A11')
   for expert in others:
    for qid in groups[expert]:
     if by[qid]['domain']!=dom or by[qid]['hazard_count']>=by[pid]['hazard_count']:continue
     new_a11=a11-by[pid]['reading_load']+by[qid]['reading_load']
     new_peer=current[expert]-by[qid]['reading_load']+by[pid]['reading_load']
     peer=[new_peer if e==expert else current[e] for e in others]
     if max(peer)-min(peer)>READING_SPREAD_LIMIT or new_a11>=min(peer):continue
     key=(by[pid]['hazard_count']-by[qid]['hazard_count'],min(peer)-new_a11,expert,qid)
     if best is None or key>best[0]:best=(key,expert,qid)
   if best is None:continue
   _,expert,qid=best
   groups['A11'].remove(pid);groups[expert].remove(qid)
   groups['A11'].append(qid);groups[expert].append(pid)
   moved=True
  if not moved:break
 for _ in range(30):
  current_h={e:hazards(e) for e in others};current_j={e:judged(e) for e in others}
  heavy=max(others,key=lambda e:(current_h[e],e));light=min(others,key=lambda e:(current_h[e],e))
  if current_h[heavy]-current_h[light]<=10:break
  best=None
  for pid in groups[heavy]:
   for qid in groups[light]:
    if by[pid]['domain']!=by[qid]['domain'] or by[pid]['hazard_count']<=by[qid]['hazard_count']:continue
    new_heavy=current_j[heavy]-by[pid]['reading_load']+by[qid]['reading_load']
    new_light=current_j[light]-by[qid]['reading_load']+by[pid]['reading_load']
    peer=[];peer_h=[]
    for expert in others:
     if expert==heavy:peer.append(new_heavy);peer_h.append(current_h[heavy]-by[pid]['hazard_count']+by[qid]['hazard_count'])
     elif expert==light:peer.append(new_light);peer_h.append(current_h[light]-by[qid]['hazard_count']+by[pid]['hazard_count'])
     else:peer.append(current_j[expert]);peer_h.append(current_h[expert])
    reduction=(current_h[heavy]-current_h[light])-(max(peer_h)-min(peer_h))
    if reduction<=0 or max(peer)-min(peer)>READING_SPREAD_LIMIT or judged('A11')>=min(peer) or hazards('A11')>min(peer_h):continue
    key=(reduction,-abs(new_heavy-new_light),pid,qid)
    if best is None or key>best[0]:best=(key,pid,qid)
  if best is None:break
  _,pid,qid=best
  groups[heavy].remove(pid);groups[light].remove(qid)
  groups[heavy].append(qid);groups[light].append(pid)
 order={p['id']:i for i,p in enumerate(index)}
 for ids in groups.values():ids.sort(key=order.get)
 assert all(len(ids)==30 for ids in groups.values())
 assert len(set(sum(groups.values(),[])))==len(index)==330
 rest=[judged(e) for e in others]
 assert judged('A11')<min(rest) and max(rest)-min(rest)<=READING_SPREAD_LIMIT
 assert sum(by[i]['hazard_count'] for i in groups['A11'])<=min(sum(by[i]['hazard_count'] for i in groups[e]) for e in others)
 return groups
def build(workbook:Path,images:Path,out:Path)->dict:
 wb=openpyxl.load_workbook(workbook,read_only=True,data_only=True)
 source_rows=list(wb.worksheets[0].iter_rows(values_only=True))
 rows=[dict(zip(source_rows[0],r)) for r in source_rows[1:] if any(v is not None for v in r)]
 image_files=[p for domain in ('營造','製造') for p in (images/domain).iterdir() if p.is_file() and p.suffix.lower() in ('.jpg','.jpeg','.png','.webp')]
 filenames=[p.name for p in image_files]
 if len(filenames)!=len(set(filenames)):raise ValueError('Duplicate image filename across domains')
 by_name={p.name:p for p in image_files}
 result_names=[r['OriginalFileName'] for r in rows]
 if len(result_names)!=len(set(result_names)):raise ValueError('Duplicate result filename')
 if set(result_names)!=set(by_name):raise ValueError(f'Result/image mismatch: {set(result_names)^set(by_name)}')
 source_hash=digest(workbook)
 image_hashes={name:digest(path) for name,path in by_name.items()}
 fingerprint=hashlib.sha256((source_hash+'\n'+'\n'.join(f'{n}:{image_hashes[n]}' for n in sorted(by_name))).encode()).hexdigest()
 dataset_id='vlm330-20261008-'+fingerprint[:12]
 parsed=[]
 for excel_row,row in enumerate(rows,start=2):
  parsed.append((excel_row,row,json.loads(row['VlmReportJson']),json.loads(row['SummaryReportJson'])))
 catalog=catalog_from_details(detail for _,_,_,summary in parsed for detail in summary.get('法規詳情') or [])
 index=[];unparsed=[]
 for excel_row,row,vlm,summary in parsed:
  name=row['OriginalFileName'];path=by_name[name];photo_id=path.stem
  details=summary.get('法規詳情',[])
  hazards=[]
  for n,item in enumerate(summary['辨識結果']):
   law_text=item.get('違反法規','') or ''
   refs,unresolved=split_refs(law_text,details,catalog)
   if not refs and not citation_present(law_text):law_text=''
   unparsed.extend({'photo_id':photo_id,'hazard_id':f'h{n+1:02d}','text':t} for t in unresolved)
   cases=[{'id':f'c{k+1:02d}','text':text,'source_status':'unconfirmed'} for k,text in enumerate(item.get('過去事故案例',[]))]
   hazards.append({'id':f'h{n+1:02d}','type':item['危害類型'],'law_text':law_text,'laws':refs,'cases':cases,
    'ppe':item.get('個人防護具',''),'improvements':item.get('建議改善方案',[]),
    'immediate_danger':item.get('立即發生危險之虞'),'source_pointer':f'/辨識結果/{n}'})
  record={'id':photo_id,'filename':name,'domain':path.parent.name,'image':'assets/photos/'+name,
   'image_sha256':image_hashes[name],'dataset_id':dataset_id,'identified_at':row['IdentifyTime'],'source_row':excel_row,
   'description':summary['整體描述'],'hazards':hazards,'overall_improvements':summary.get('整體改善建議',[]),
   'vlm':vlm.get('HazardItems',[]),'ai_law_details':details}
  write_json(out/'data/cases'/f'{photo_id}.json',record)
  image_out=out/'assets/photos'/name
  image_out.parent.mkdir(parents=True,exist_ok=True)
  if not image_out.exists() or digest(image_out)!=image_hashes[name]:shutil.copy2(path,image_out)
  index.append({'id':photo_id,'filename':name,'domain':path.parent.name,'image':record['image'],'image_sha256':image_hashes[name],
   'data_file':f'data/cases/{photo_id}.json','hazard_count':len(hazards),'hazard_types':[h['type'] for h in hazards],
   'case_count':sum(len(h['cases']) for h in hazards),'reading_load':reading_load(record)})
 index.sort(key=lambda r:(r['filename'][0],int(re.search(r'\d+',r['filename']).group()),r['filename']))
 assignments=assign_photos(index)
 assignment_version='disjoint30-'+hashlib.sha256(json.dumps(assignments,sort_keys=True,separators=(',',':')).encode()).hexdigest()[:12]
 manifest={'schema_version':2,'id':dataset_id,'title':'照片危害辨識專家審閱','test_date':'2026-10-08',
  'source_workbook':workbook.name,'source_workbook_sha256':source_hash,'fingerprint':fingerprint,
  'expert_ids':[f'A{i:02d}' for i in range(1,12)],'assignment':'disjoint_30_per_expert',
  'assignment_version':assignment_version,'assignments':assignments,'photos_per_expert':30,
  'photo_count':len(index),'domains':dict(Counter(r['domain'] for r in index)),
  'hazard_count':sum(r['hazard_count'] for r in index),'case_reference_count':sum(r['case_count'] for r in index),
  'hazard_types':HAZARD_TYPES,'law_baseline_date':'2026-10-08','case_sources':'unconfirmed',
  'photos':index}
 write_json(out/'data/manifest.json',manifest)
 write_assignments(manifest,out)
 write_json(out/'docs/import-report.json',{'dataset_id':dataset_id,'photo_count':len(index),'filename_matches':len(index),
  'hazard_count':manifest['hazard_count'],'case_reference_count':manifest['case_reference_count'],
  'unparsed_law_fragments':unparsed,'photo_copy_mode':'Original bytes, no crop, rotation or recompression',
  'metadata_source':'Actual revised photo files; stale index and manifest not used'})
 return {k:v for k,v in manifest.items() if k not in ('photos','hazard_types')}
def write_assignments(manifest:dict,out:Path)->None:
 photos={p['id']:p for p in manifest['photos']}
 lines=['# 固定分組：每位 30 張','',
  f"分組版本：{manifest['assignment_version']}。依待評閱讀量分配（整體描述、情境理由、法規條文、職災案例）。A01–A10 彼此相差在 800 字以內，A11 閱讀量最少。營造／製造仍是十組 20/10、A11 為 19/11。十一組互不重複，涵蓋 330 張。",'']
 for expert in manifest['expert_ids']:
  items=[photos[i] for i in manifest['assignments'][expert]]
  cons=sum(p['domain']=='營造' for p in items);mfg=sum(p['domain']=='製造' for p in items)
  lines+=[f'## {expert}','',
   f"30 張：營造 {cons} 張、製造 {mfg} 張；待評文字 {sum(p['reading_load'] for p in items):,} 字，危害 {sum(p['hazard_count'] for p in items)} 項。",'',
   f'[專家入口](https://fefe9487.github.io/vlm330-expert-review/?expert={expert})','',
   '、'.join(p['filename'] for p in items),'']
 (out/'docs/ASSIGNMENTS.md').write_text('\n'.join(lines),encoding='utf-8')
def reassign(out:Path)->dict:
 manifest=json.loads((out/'data/manifest.json').read_text(encoding='utf-8'))
 for photo in manifest['photos']:
  photo['reading_load']=reading_load(json.loads((out/photo['data_file']).read_text(encoding='utf-8')))
 manifest['assignments']=assign_photos(manifest['photos'])
 manifest['assignment_version']='disjoint30-'+hashlib.sha256(json.dumps(manifest['assignments'],sort_keys=True,separators=(',',':')).encode()).hexdigest()[:12]
 write_json(out/'data/manifest.json',manifest)
 write_assignments(manifest,out)
 photos={p['id']:p for p in manifest['photos']}
 return {e:{'reading_load':sum(photos[i]['reading_load'] for i in ids),'hazards':sum(photos[i]['hazard_count'] for i in ids),
  'construction':sum(photos[i]['domain']=='營造' for i in ids),'manufacturing':sum(photos[i]['domain']=='製造' for i in ids)}
  for e,ids in manifest['assignments'].items()}
def check_parser()->None:
 details=[{'法規名稱':'職業安全衛生設施規則','條號':'第155-1條','完整法規內文':'吊運條文'},
  {'法規名稱':'職業安全衛生設施規則','條號':'第116條','完整法規內文':'車輛'},
  {'法規名稱':'職業安全衛生設施規則','條號':'第119條','完整法規內文':'營建機械'},
  {'法規名稱':'營造安全衛生設施標準','條號':'第69條','完整法規內文':'開挖'},
  {'法規名稱':'營造安全衛生設施標準','條號':'第159條','完整法規內文':'拆除'},
  {'法規名稱':'營造安全衛生設施標準第17條','條號':'第17條','完整法規內文':'假設'}]
 refs,missed=split_refs('職業安全衛生設施規則第155條之1',details,{})
 assert not missed and refs[0]['article']=='155-1' and refs[0]['ai_texts']==['吊運條文'] and refs[0]['official_url'].endswith('flno=155-1')
 refs,missed=split_refs('職業安全衛生設施規則第116條、第119條；營造安全衛生設施標準第69條、第159條；職業安全衛生法第32條',details,{('職業安全衛生法','32'):'教育訓練'})
 assert not missed and [(r['law_name'],r['article']) for r in refs]==[
  ('職業安全衛生設施規則','116'),('職業安全衛生設施規則','119'),('營造安全衛生設施標準','69'),('營造安全衛生設施標準','159'),('職業安全衛生法','32')]
 assert refs[4]['ai_texts']==['教育訓練'] and 'flno=32' in refs[4]['official_url']
 refs,_=split_refs('職業安全衛生設施規則第29條之1、第184條之1、第196條',[],{})
 assert [r['article'] for r in refs]==['29-1','184-1','196']
 refs,_=split_refs('營造安全衛生設施規則第224條',[{'法規名稱':'營造安全衛生設施標準','條號':'第224條','完整法規內文':'高處'}],{})
 assert refs[0]['law_name']=='營造安全衛生設施標準' and refs[0]['ai_texts']==['高處'] and 'N0060014' in refs[0]['official_url']
 refs,_=split_refs('無指定法規',details,{});assert refs==[]
 refs,_=split_refs('高氣溫作業熱危害預防指引第4條',[{'法規名稱':'高氣溫作業熱危害預防指引','條號':'第4條','完整法規內文':'熱'}],{})
 assert refs[0]['ai_texts']==['熱'] and refs[0]['official_url'].startswith('https://www.osha.gov.tw/')
 assert detail_identity(details[-1])==('營造安全衛生設施標準','17')
def repair_laws(out:Path,write:bool=True)->dict:
 """Rewrite citation cards in place. Photos, hazard judgments and expert groups stay as published."""
 check_parser()
 paths=sorted((out/'data/cases').glob('*.json'))
 cases=[json.loads(path.read_text(encoding='utf-8')) for path in paths]
 catalog=catalog_from_details(detail for case in cases for detail in case.get('ai_law_details') or [])
 stats=Counter();samples=[]
 for case in cases:
  for hazard in case.get('hazards') or []:
   original=hazard.get('law_text') or ''
   refs,unresolved=split_refs(original,case.get('ai_law_details') or [],catalog)
   if not refs and not citation_present(original):
    if original.strip():stats['cleared_placeholders']+=1
    hazard['law_text']=''
   hazard['laws']=refs
   stats['law_cards']+=len(refs)
   stats['cards_without_text']+=sum(not ref['ai_texts'] for ref in refs)
   stats['unresolved']+=len(unresolved)
   if unresolved and len(samples)<12:samples.append({'photo':case['id'],'hazard':hazard['id'],'text':original,'unresolved':unresolved})
  stats['hazards']+=len(case.get('hazards') or [])
 if write:
  for path,case in zip(paths,cases):write_json(path,case)
  write_json(out/'docs/law-display-repair.json',{'cleared_placeholders':stats['cleared_placeholders'],'law_cards':stats['law_cards'],
   'cards_without_text':stats['cards_without_text'],'unresolved':stats['unresolved'],'samples':samples,
   'note':'Citation display only. Hazard identification and expert assignments were not changed.'})
 return {'photos':len(cases),'hazards':stats['hazards'],'law_cards':stats['law_cards'],'cards_without_text':stats['cards_without_text'],
  'cleared_placeholders':stats['cleared_placeholders'],'unresolved':stats['unresolved'],'samples':samples}
if __name__=='__main__':
 sys.stdout.reconfigure(encoding='utf-8')
 parser=argparse.ArgumentParser()
 parser.add_argument('--workbook',type=Path,default=DEFAULT_PROJECT.parent/'sh168圖形辨識330筆結果.xlsx')
 parser.add_argument('--images',type=Path,default=DEFAULT_PROJECT.parents[1]/'危害辨識測試_330')
 parser.add_argument('--output',type=Path,default=DEFAULT_PROJECT)
 parser.add_argument('--reassign',action='store_true')
 parser.add_argument('--repair-laws',action='store_true')
 parser.add_argument('--dry-run',action='store_true')
 args=parser.parse_args()
 if args.repair_laws:
  print(json.dumps(repair_laws(args.output,write=not args.dry_run),ensure_ascii=False,indent=2))
 elif args.reassign:
  loads=reassign(args.output)
  version=json.loads((args.output/'data/manifest.json').read_text(encoding='utf-8'))['assignment_version']
  print(json.dumps({'assignment_version':version,'loads':loads},ensure_ascii=False,indent=2))
 else:print(json.dumps(build(args.workbook,args.images,args.output),ensure_ascii=False,indent=2))
